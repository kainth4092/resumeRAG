import docx

from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_items(document):
    """
    Yield DOCX paragraphs and tables in their original document order.
    """

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(
                child,
                document,
            )

        elif child.tag.endswith("}tbl"):
            yield Table(
                child,
                document,
            )


def extract_table_rows(table: Table) -> list[str]:
    """
    Extract non-empty table cells while preserving row order.

    Duplicate cell text inside the same row is removed because merged
    DOCX cells can sometimes expose the same content more than once.
    """

    output = []

    for row in table.rows:
        row_values = []
        seen = set()

        for cell in row.cells:
            cell_text = " ".join(cell.text.split())

            if not cell_text:
                continue

            normalized = cell_text.casefold()

            if normalized in seen:
                continue

            seen.add(normalized)
            row_values.append(cell_text)

        if row_values:
            output.append(" | ".join(row_values))

    return output


def extract_text_from_docx(
    file_path: str,
) -> str:
    """
    Extract paragraphs, tables, headers, footers, and textboxes in DOCX.
    Preserving order is required for section-based resume parsing.
    Headers, footers, and textboxes frequently contain contact info and personal links.
    """
    document = docx.Document(file_path)

    header_text = []
    footer_text = []
    for section in document.sections:
        try:
            for p in section.header.paragraphs:
                t = p.text.strip()
                if t:
                    header_text.append(t)
        except Exception:
            pass
        try:
            for p in section.footer.paragraphs:
                t = p.text.strip()
                if t:
                    footer_text.append(t)
        except Exception:
            pass

    body_text = []
    for block in iter_block_items(document):
        if isinstance(
            block,
            Paragraph,
        ):
            paragraph_text = block.text.strip()
            if paragraph_text:
                body_text.append(paragraph_text)

        elif isinstance(
            block,
            Table,
        ):
            body_text.extend(extract_table_rows(block))

    # Extract text from text boxes (w:txbxContent)
    textbox_text = []
    try:
        for txbx in document.element.xpath(".//w:txbxContent"):
            for p in txbx.xpath(".//w:p"):
                p_obj = Paragraph(p, document)
                t = p_obj.text.strip()
                if t and t not in body_text:
                    textbox_text.append(t)
    except Exception:
        pass

    all_parts = []
    if header_text:
        all_parts.extend(header_text)
    if textbox_text:
        all_parts.extend(textbox_text)
    if body_text:
        all_parts.extend(body_text)
    if footer_text:
        all_parts.extend(footer_text)

    return "\n".join(all_parts)
